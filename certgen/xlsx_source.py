"""Spreadsheet reading.

read_records_from_xlsx() is the baseline behavior, ported unchanged:
positional columns A=Nome, B=CPF, C=Turma, D=Data, E=Email, no header
names recognized. Kept as the default so nothing that already works
breaks.

Phase 3 adds an explicit column-mapping path (sniff_headers /
detect_default_mapping / read_records_with_mapping) so a spreadsheet
whose columns aren't in that exact order doesn't just silently misread
names into the CPF field -- that was the single biggest source of
confusion the plan flagged for non-technical users.
"""
import datetime
import unicodedata
from typing import Dict, List, Optional

from openpyxl import load_workbook

from .fields import FIELD_CPF, FIELD_DATE, FIELD_EMAIL, FIELD_NAME, FIELD_TURMA

MAPPABLE_FIELDS = [FIELD_NAME, FIELD_CPF, FIELD_TURMA, FIELD_DATE, FIELD_EMAIL]
_MAPPABLE_FIELDS = MAPPABLE_FIELDS  # backward-compat alias

_HEADER_ALIASES = {
    FIELD_NAME: {"nome", "name", "participante", "aluno"},
    FIELD_CPF: {"cpf", "documento"},
    FIELD_TURMA: {"turma", "classe", "class", "grupo"},
    FIELD_DATE: {"data", "date", "data de conclusao", "data de emissao"},
    FIELD_EMAIL: {"email", "e-mail", "mail"},
}


def _normalize_header(text: str) -> str:
    text = (text or "").strip().lower()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return text


def _format_date_cell(datev) -> str:
    if isinstance(datev, datetime.date):
        return datev.strftime("%d/%m/%Y")
    if datev is None or str(datev).strip() == "":
        return ""
    return str(datev).strip()


def read_records_from_xlsx(path: str, skip_header: bool = False) -> List[Dict[str, str]]:
    """
    XLSX: Column A = Nome (required)
          Column B = CPF (optional)
          Column C = Turma (optional)
          Column D = Data (optional; empty if absent)
          Column E = Email (required only if emailing)
    skip_header=True skips the first row. Reads the active worksheet.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    records: List[Dict[str, str]] = []

    rows = ws.iter_rows(values_only=True)
    if skip_header:
        try:
            next(rows)
        except StopIteration:
            wb.close()
            return records

    for row in rows:
        a = row[0] if len(row) >= 1 else None
        if a is None or str(a).strip() == "":
            continue

        name = str(a).strip()
        cpf = str(row[1]).strip() if (len(row) >= 2 and row[1] is not None) else ""
        turma = str(row[2]).strip() if (len(row) >= 3 and row[2] is not None) else ""
        date_str = _format_date_cell(row[3] if len(row) >= 4 else None)
        email = str(row[4]).strip() if (len(row) >= 5 and row[4] is not None) else ""

        records.append({
            FIELD_NAME: name,
            FIELD_CPF: cpf,
            FIELD_TURMA: turma,
            FIELD_DATE: date_str,
            FIELD_EMAIL: email,
        })

    wb.close()
    return records


def sniff_headers(path: str) -> List[str]:
    """First row's raw cell values, for a column-mapping UI to show."""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    try:
        first_row = next(ws.iter_rows(values_only=True))
    except StopIteration:
        first_row = ()
    wb.close()
    return [("" if v is None else str(v).strip()) for v in first_row]


def detect_default_mapping(headers: List[str]) -> Dict[str, Optional[int]]:
    """Best-effort guess of {field: column_index} from header text. Falls
    back to the classic positional layout (A-E) for any field it can't
    match by name, so a spreadsheet with no recognizable headers behaves
    exactly like the original positional reader."""
    mapping: Dict[str, Optional[int]] = {}
    normalized = [_normalize_header(h) for h in headers]
    used_indices = set()

    for field_key in _MAPPABLE_FIELDS:
        found = None
        for idx, header in enumerate(normalized):
            if header in _HEADER_ALIASES[field_key] and idx not in used_indices:
                found = idx
                break
        mapping[field_key] = found
        if found is not None:
            used_indices.add(found)

    # positional fallback for anything not matched by name -- but never
    # steal a column index another field already claimed by header name.
    positional = [FIELD_NAME, FIELD_CPF, FIELD_TURMA, FIELD_DATE, FIELD_EMAIL]
    for i, field_key in enumerate(positional):
        if mapping[field_key] is None and i < len(headers) and i not in used_indices:
            mapping[field_key] = i
            used_indices.add(i)

    return mapping


def read_records_with_mapping(path: str, mapping: Dict[str, Optional[int]],
                               skip_header: bool = True) -> List[Dict[str, str]]:
    """Read records using an explicit {field: column_index or None} map.
    Any field mapped to None is left blank for every record."""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    records: List[Dict[str, str]] = []

    rows = ws.iter_rows(values_only=True)
    if skip_header:
        try:
            next(rows)
        except StopIteration:
            wb.close()
            return records

    name_col = mapping.get(FIELD_NAME)

    def cell(row, idx: Optional[int]):
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    for row in rows:
        name_val = cell(row, name_col)
        if name_val is None or str(name_val).strip() == "":
            continue

        cpf_val = cell(row, mapping.get(FIELD_CPF))
        turma_val = cell(row, mapping.get(FIELD_TURMA))
        date_val = cell(row, mapping.get(FIELD_DATE))
        email_val = cell(row, mapping.get(FIELD_EMAIL))

        records.append({
            FIELD_NAME: str(name_val).strip(),
            FIELD_CPF: str(cpf_val).strip() if cpf_val is not None else "",
            FIELD_TURMA: str(turma_val).strip() if turma_val is not None else "",
            FIELD_DATE: _format_date_cell(date_val),
            FIELD_EMAIL: str(email_val).strip() if email_val is not None else "",
        })

    wb.close()
    return records
