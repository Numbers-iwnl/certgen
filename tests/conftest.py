"""Shared fixtures for the certgen test suite."""
import datetime
import glob
import os
import sys

import pymupdf as fitz
import pytest
from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)


def _find_system_font(*names: str) -> str:
    candidates = []
    windir = os.environ.get("WINDIR", r"C:\Windows")
    for name in names:
        candidates.append(os.path.join(windir, "Fonts", name))
    for c in candidates:
        if os.path.isfile(c):
            return c
    for pattern in ("/usr/share/fonts/**/*.ttf", "/System/Library/Fonts/*.ttf"):
        found = glob.glob(pattern, recursive=True)
        if found:
            return found[0]
    pytest.skip("no usable .ttf font found on this system")


@pytest.fixture(scope="session")
def font_regular() -> str:
    return _find_system_font("arial.ttf", "DejaVuSans.ttf")


@pytest.fixture(scope="session")
def font_bold() -> str:
    return _find_system_font("arialbd.ttf", "DejaVuSans-Bold.ttf")


@pytest.fixture
def blank_pdf(tmp_path):
    path = str(tmp_path / "blank.pdf")
    doc = fitz.open()
    doc.new_page(width=595, height=842)
    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def certificate_pdf(tmp_path, font_regular, font_bold):
    path = str(tmp_path / "certificate.pdf")
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)

    page.insert_text((300, 300), "(Seu Nome Aqui)", fontsize=24,
                      fontfile=font_regular, fontname="F0")
    page.insert_text((80, 500), "CPF:", fontsize=12,
                      fontfile=font_regular, fontname="F1")
    page.insert_text((300, 500), "TURMA:", fontsize=12,
                      fontfile=font_bold, fontname="TurmaBold")
    page.insert_text((500, 500), "DATA:", fontsize=12,
                      fontfile=font_regular, fontname="F2")

    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def certificate_pdf_date_slashes(tmp_path, font_regular):
    path = str(tmp_path / "certificate_slashes.pdf")
    doc = fitz.open()
    page = doc.new_page(width=842, height=595)
    page.insert_text((300, 300), "(Seu Nome Aqui)", fontsize=24,
                      fontfile=font_regular, fontname="F0")
    page.insert_text((400, 520), "01/01/2026", fontsize=12,
                      fontfile=font_regular, fontname="F1")
    page.insert_text((550, 520), "31/12/2026", fontsize=12,
                      fontfile=font_regular, fontname="F2")
    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def sample_xlsx_no_header(tmp_path):
    path = str(tmp_path / "sample_no_header.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["Ana Beatriz", "12345678900", "Turma A", datetime.date(2026, 3, 15), "ana@example.com"])
    ws.append(["Bruno Costa", "", "", "", ""])
    ws.append(["", "", "", "", ""])
    ws.append(["Carla Dias com Nome Bem Comprido De Verdade", "98765432100", "Turma B", "10/02/2026", "carla@example.com"])
    wb.save(path)
    return path


@pytest.fixture
def sample_xlsx_with_header(tmp_path):
    path = str(tmp_path / "sample_with_header.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["Nome", "CPF", "Turma", "Data", "Email"])
    ws.append(["Diego Alves", "11122233344", "Turma C", datetime.date(2026, 6, 1), "diego@example.com"])
    wb.save(path)
    return path


@pytest.fixture
def sample_xlsx_reordered_headers(tmp_path):
    """Columns not in the classic A-E order, to exercise column mapping."""
    path = str(tmp_path / "sample_reordered.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["Email", "Nome", "Turma"])
    ws.append(["elis@example.com", "Elis Regina", "Turma D"])
    wb.save(path)
    return path


@pytest.fixture(autouse=True)
def isolate_appdata(tmp_path, monkeypatch):
    """Every test gets its own fake %APPDATA% so prefs/presets/logs never
    leak between tests or touch the real user profile."""
    fake_appdata = tmp_path / "AppData_Roaming"
    fake_appdata.mkdir()
    monkeypatch.setenv("APPDATA", str(fake_appdata))
    yield
